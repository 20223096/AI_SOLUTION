#!/usr/bin/env python3
"""Streaming ETL for supplied S-DoT CSV and Sinchon road-noise workbooks."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import statistics
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

TZ = ZoneInfo("Asia/Seoul")
SDOT_REQUIRED = {"model": (0, "모델명"), "serial": (1, "시리얼"), "observed": (2, "센서 시간"), "district": (4, "자치구역"), "neighborhood": (5, "행정구역"), "max": (24, "최대 소음"), "avg": (25, "평균 소음"), "min": (26, "최소 소음"), "collected": (57, "데이터수집시간"), "type": (58, "데이터구분번호")}
SDOT_FIELDS = ["sensor_id", "sensor_serial", "observed_at", "district", "neighborhood", "noise_avg", "noise_max", "noise_min", "record_type", "collected_at", "quality_status", "quality_reasons", "duplicate_status", "source_file", "source_sha256", "provenance"]
ROAD_FIELDS = ["measurement_date", "hour_ending", "station_name", "noise_level", "metric", "source_period", "source_file", "source_member", "quality_status", "provenance"]


def sha256_file(path: Path, block_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(block_size), b""):
            digest.update(block)
    return digest.hexdigest()


def parse_sensor_time(value: str) -> datetime:
    return datetime.strptime(value.strip(), "%Y-%m-%d_%H:%M:%S").replace(tzinfo=TZ)


def _number(value: str) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def normalize_sdot_row(row: list[str], source_file: str, file_hash: str) -> dict:
    observed = parse_sensor_time(row[2])
    low, avg, high = _number(row[26]), _number(row[25]), _number(row[24])
    reasons = []
    if None in (low, avg, high): reasons.append("missing_noise_value")
    if any(value is not None and value < 0 for value in (low, avg, high)): reasons.append("negative_noise_value")
    if None not in (low, avg, high) and not low <= avg <= high: reasons.append("min_avg_max_order_violation")
    return {
        "sensor_id": row[0], "sensor_serial": row[1], "observed_at": observed.isoformat(),
        "district": row[4], "neighborhood": row[5], "noise_avg": avg, "noise_max": high, "noise_min": low,
        "record_type": row[58], "collected_at": row[57], "quality_status": "valid" if not reasons else "held",
        "quality_reasons": "|".join(reasons), "duplicate_status": "unique", "source_file": source_file,
        "source_sha256": file_hash,
        "provenance": json.dumps({"status":"verified_public", "dataset":"S-DoT environmental information", "raw_observed_at":row[2], "metric":"average/max/min noise; exact acoustic definition not supplied"}, ensure_ascii=False),
    }


def validate_sdot_header(header: list[str]) -> None:
    problems = [f"{name}@{index}" for index, name in (item for item in SDOT_REQUIRED.values()) if len(header) <= index or header[index] != name]
    if problems:
        raise ValueError("S-DoT header mismatch: " + ", ".join(problems))


def process_sdot(paths: list[Path], output_path: Path) -> tuple[dict, list[dict]]:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    hashes, seen_hashes, header_signatures = {}, {}, Counter()
    total_rows = extracted_rows = duplicate_rows = conflict_rows = 0
    missing = Counter(); abnormal = Counter(); sensors = Counter(); record_types = Counter()
    logical_seen: dict[tuple[str, str], tuple] = {}
    model_time_seen=set(); collected_seen=set(); type_collected_seen=set()
    alternate_duplicates=Counter()
    sensor_times = defaultdict(list); valid_records = []
    with output_path.open("w", encoding="utf-8", newline="") as out:
        writer = csv.DictWriter(out, fieldnames=SDOT_FIELDS); writer.writeheader()
        for path in paths:
            digest = sha256_file(path); hashes[path.name] = digest
            if digest in seen_hashes:
                continue
            seen_hashes[digest] = path.name
            with path.open(encoding="cp949", newline="") as stream:
                reader = csv.reader(stream); header = next(reader); validate_sdot_header(header)
                header_signatures[hashlib.sha256(json.dumps(header, ensure_ascii=False).encode()).hexdigest()] += 1
                for row in reader:
                    total_rows += 1
                    if len(row) < 59 or row[4] != "Seodaemun-gu" or row[5] != "Sinchon-dong":
                        continue
                    extracted_rows += 1
                    try:
                        item = normalize_sdot_row(row, path.name, digest)
                    except (ValueError, IndexError):
                        abnormal["unparseable_row"] += 1
                        item = {field: None for field in SDOT_FIELDS}
                        item.update({"sensor_id":row[0] if row else None,"sensor_serial":row[1] if len(row)>1 else None,"district":row[4] if len(row)>4 else None,"neighborhood":row[5] if len(row)>5 else None,"quality_status":"held","quality_reasons":"unparseable_row","duplicate_status":"not_evaluated","source_file":path.name,"source_sha256":digest,"provenance":json.dumps({"status":"verified_public","dataset":"S-DoT environmental information","raw_row_preserved":row},ensure_ascii=False)})
                        writer.writerow(item); continue
                    key = (item["sensor_serial"], item["observed_at"])
                    model_key=(item["sensor_id"],item["observed_at"]); collected_key=item["collected_at"]; type_collected_key=(item["record_type"],item["collected_at"])
                    if model_key in model_time_seen: alternate_duplicates["sensor_id_plus_time"] += 1
                    else: model_time_seen.add(model_key)
                    if collected_key and collected_key in collected_seen: alternate_duplicates["collection_time"] += 1
                    elif collected_key: collected_seen.add(collected_key)
                    if collected_key and type_collected_key in type_collected_seen: alternate_duplicates["record_type_plus_collection_time"] += 1
                    elif collected_key: type_collected_seen.add(type_collected_key)
                    signature = (item["noise_avg"], item["noise_max"], item["noise_min"], item["record_type"], item["collected_at"])
                    if key in logical_seen:
                        duplicate_rows += 1
                        if logical_seen[key] != signature:
                            conflict_rows += 1; item["duplicate_status"] = "conflicting_duplicate_held"
                        else:
                            item["duplicate_status"] = "exact_logical_duplicate_excluded"
                    else:
                        logical_seen[key] = signature
                    for field in ("noise_avg", "noise_max", "noise_min"):
                        if item[field] is None: missing[field] += 1
                    if item["quality_status"] != "valid":
                        for reason in item["quality_reasons"].split("|"): abnormal[reason] += 1
                    sensors[item["sensor_serial"]] += 1; record_types[item["record_type"]] += 1
                    writer.writerow(item)
                    if item["quality_status"] == "valid" and item["duplicate_status"] == "unique":
                        dt = datetime.fromisoformat(item["observed_at"]); sensor_times[item["sensor_serial"]].append(dt); valid_records.append(item)
    intervals = Counter()
    for values in sensor_times.values():
        values.sort()
        intervals.update(int((b-a).total_seconds()/60) for a,b in zip(values, values[1:]) if b > a)
    times = [datetime.fromisoformat(x["observed_at"]) for x in valid_records]
    possible_sensor_hours = 0
    observed_sensor_hours = len({(x["sensor_serial"], datetime.fromisoformat(x["observed_at"]).replace(minute=0, second=0)) for x in valid_records})
    if times and sensors:
        possible_sensor_hours = len(sensors) * (int((max(times)-min(times)).total_seconds()/3600)+1)
    report = {
        "files": len(paths), "unique_file_hashes": len(seen_hashes), "duplicate_files": len(paths)-len(seen_hashes), "hashes": hashes,
        "header_signature_count": len(header_signatures), "total_rows": total_rows, "sinchon_rows": extracted_rows,
        "valid_unique_rows": len(valid_records), "sensor_count": len(sensors), "sensor_rows": dict(sensors),
        "period_start": min((x.isoformat() for x in times), default=None), "period_end": max((x.isoformat() for x in times), default=None),
        "missing_counts": dict(missing), "abnormal_counts": dict(abnormal), "logical_duplicate_rows": duplicate_rows,
        "conflicting_duplicate_rows": conflict_rows, "record_types": dict(record_types),
        "alternate_duplicate_key_counts": dict(alternate_duplicates),
        "measurement_interval_minutes_top": intervals.most_common(10), "hourly_availability_ratio": round(observed_sensor_hours/possible_sensor_hours, 6) if possible_sensor_hours else None,
    }
    return report, valid_records


def _road_rows_from_workbook(workbook, source_file: str, source_member: str) -> list[dict]:
    sheet_name = next((name for name in workbook.sheetnames if "신촌" in name), None)
    if not sheet_name: return []
    ws = workbook[sheet_name]; rows = []
    if ws.max_column >= 28 and ws.cell(4, 28).value == "LEQ":
        metric = "LEQ"; station = str(ws.cell(4, 15).value or "신촌")
        for values in ws.iter_rows(min_row=7, values_only=True):
            day = values[1]
            if not isinstance(day, datetime): continue
            for hour in range(1, 25):
                level = values[hour+1]
                if not isinstance(level, (int, float)): continue
                rows.append({"measurement_date":day.date().isoformat(),"hour_ending":hour,"station_name":station,"noise_level":float(level),"metric":metric,"source_period":day.strftime("%Y-%m"),"source_file":source_file,"source_member":source_member,"quality_status":"valid" if level >= 0 else "held_negative","provenance":json.dumps({"status":"verified_public","sheet":sheet_name,"measurement_place":station,"address":None,"coordinates":None,"metric_label":"LEQ"},ensure_ascii=False)})
    elif ws.cell(1, 1).value == "년도":
        metric = "소음도_정의미제공"
        for values in ws.iter_rows(min_row=2, values_only=True):
            if not all(isinstance(values[i], (int, float)) for i in (0,1,2,3,5)): continue
            day = datetime(int(values[0]),int(values[1]),int(values[2]))
            rows.append({"measurement_date":day.date().isoformat(),"hour_ending":int(values[3]),"station_name":str(values[4]),"noise_level":float(values[5]),"metric":metric,"source_period":day.strftime("%Y-%m"),"source_file":source_file,"source_member":source_member,"quality_status":"valid" if values[5] >= 0 else "held_negative","provenance":json.dumps({"status":"verified_public","sheet":sheet_name,"address":None,"coordinates":None,"metric_label":"소음도; definition absent"},ensure_ascii=False)})
    return rows


def process_road(paths: list[Path], output_path: Path) -> tuple[dict, list[dict]]:
    output_path.parent.mkdir(parents=True, exist_ok=True); rows=[]; inventory=[]
    for path in paths:
        digest=sha256_file(path); entry={"file":path.name,"size":path.stat().st_size,"sha256":digest,"members":[]}
        if path.suffix.lower()==".zip":
            with zipfile.ZipFile(path) as archive:
                for member in archive.namelist():
                    if not member.lower().endswith(".xlsx"): continue
                    entry["members"].append(member)
                    workbook=openpyxl.load_workbook(io.BytesIO(archive.read(member)),read_only=True,data_only=True)
                    rows.extend(_road_rows_from_workbook(workbook,path.name,member))
        elif path.suffix.lower()==".xlsx":
            workbook=openpyxl.load_workbook(path,read_only=True,data_only=True)
            rows.extend(_road_rows_from_workbook(workbook,path.name,path.name))
        inventory.append(entry)
    with output_path.open("w",encoding="utf-8",newline="") as out:
        writer=csv.DictWriter(out,fieldnames=ROAD_FIELDS); writer.writeheader(); writer.writerows(rows)
    valid=[x for x in rows if x["quality_status"]=="valid"]
    report={"files":len(paths),"inventory":inventory,"rows":len(rows),"valid_rows":len(valid),"stations":dict(Counter(x["station_name"] for x in rows)),"metrics":dict(Counter(x["metric"] for x in rows)),"period_start":min((x["measurement_date"] for x in rows),default=None),"period_end":max((x["measurement_date"] for x in rows),default=None),"negative_rows":sum(x["quality_status"]!="valid" for x in rows),"location_metadata":{"address":None,"coordinates":None,"warning":"supplied workbooks identify Sinchon by sheet/place name only"}}
    return report, valid


def build_profiles(sdot: list[dict], road: list[dict], output_path: Path) -> tuple[list[dict], dict]:
    buckets=defaultdict(list)
    for row in sdot:
        dt=datetime.fromisoformat(row["observed_at"]); buckets[("sdot_observed_profile","weekend" if dt.weekday()>=5 else "weekday",dt.hour)].append(row["noise_avg"])
    for row in road:
        if row["metric"] != "LEQ" or not ("2025-01" <= row["source_period"] <= "2026-06"): continue
        dt=datetime.fromisoformat(row["measurement_date"]); buckets[("road_traffic_observed_profile","weekend" if dt.weekday()>=5 else "weekday",row["hour_ending"]%24)].append(row["noise_level"])
    output_path.parent.mkdir(parents=True,exist_ok=True)
    rows=[]
    for (dataset,day_type,hour),values in sorted(buckets.items()):
        rows.append({"profile_id":dataset,"day_type":day_type,"hour":hour,"representative_db":round(statistics.median(values),2),"mean_db":round(statistics.fmean(values),2),"observation_count":len(values),"metric":"LEQ" if dataset.startswith("road") else "device_average_noise_metric_unspecified","calculation_eligible":dataset.startswith("road")})
    with output_path.open("w",encoding="utf-8",newline="") as out:
        writer=csv.DictWriter(out,fieldnames=list(rows[0]) if rows else ["profile_id"]); writer.writeheader(); writer.writerows(rows)
    return rows, {"profile_rows":len(rows),"method":"median by dataset/day_type/hour; S-DoT and road LEQ never combined"}


def write_backend_catalog(profile_rows: list[dict], sdot_report: dict, road_report: dict, target: Path) -> None:
    def values(profile): return [{"day_type":x["day_type"],"hour":x["hour"],"representative_db":x["representative_db"],"observation_count":x["observation_count"]} for x in profile_rows if x["profile_id"]==profile]
    catalog=[
      {"id":"assumed_demo","label":"기존 PoC 가정 배경","calculation_eligible":True,"metric":"LAeq_dBA_assumed","observation_period":None,"hourly_values":[],"source_metadata":{"status":"assumed","source":"PoC scenario settings","unit":"LAeq_dBA"},"quality_warnings":["실측값이 아닌 기존 시연 가정값"]},
      {"id":"sdot_observed_profile","label":"S-DoT 신촌 시간 패턴","calculation_eligible":False,"metric":"device_average_noise_metric_unspecified","observation_period":{"start":sdot_report["period_start"],"end":sdot_report["period_end"]},"hourly_values":values("sdot_observed_profile"),"source_metadata":{"status":"verified_public","source":"Seoul S-DoT environmental CSV","unit":"source noise unit not documented in supplied files"},"quality_warnings":["측정지표·교정·설치 위치 호환성이 확인되지 않아 물리 합성에 사용하지 않음"]},
      {"id":"road_traffic_observed_profile","label":"신촌 도로교통소음 시간별 LEQ 대표 패턴","calculation_eligible":True,"metric":"LEQ","observation_period":{"start":"2025-01-01","end":"2026-06-30"},"hourly_values":values("road_traffic_observed_profile"),"source_metadata":{"status":"verified_public","source":"Seoul road traffic noise monthly workbooks: Sinchon hourly sheets","unit":"dB LEQ"},"quality_warnings":["시뮬레이션 날짜의 실측이 아닌 과거 시간대별 중앙값","주소·좌표가 없어 다른 수음점의 실측값으로 간주할 수 없음","명시적 선택 시 지역 대표 배경 후보로만 사용"]}
    ]
    target.write_text(json.dumps(catalog,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")


def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--raw-root",type=Path,default=Path("data/raw")); parser.add_argument("--processed-root",type=Path,default=Path("data/processed")); parser.add_argument("--backend-catalog",type=Path,default=Path("backend/app/data/background_profiles.json")); args=parser.parse_args()
    sdot_paths=sorted((args.raw_root/"sdot").glob("*.csv")); road_paths=sorted((args.raw_root/"road_traffic_noise").glob("*"))
    sdot_report,sdot_rows=process_sdot(sdot_paths,args.processed_root/"sdot_sinchon_clean.csv")
    road_report,road_rows=process_road(road_paths,args.processed_root/"road_noise_sinchon_clean.csv")
    profile_rows,profile_report=build_profiles(sdot_rows,road_rows,args.processed_root/"background_profile.csv")
    write_backend_catalog(profile_rows,sdot_report,road_report,args.backend_catalog)
    report={"generated_at":datetime.now(TZ).isoformat(),"sdot":sdot_report,"road":road_report,"profiles":profile_report}
    (args.processed_root/"etl_report.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"sdot":{k:sdot_report[k] for k in ("files","unique_file_hashes","total_rows","sinchon_rows","valid_unique_rows","sensor_count","logical_duplicate_rows","conflicting_duplicate_rows")},"road":{k:road_report[k] for k in ("files","rows","valid_rows","metrics")},"profiles":profile_report},ensure_ascii=False,indent=2))


if __name__=="__main__": main()
